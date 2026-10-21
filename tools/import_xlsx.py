#!/usr/bin/env python3
"""Convert the training workbook (xlsx) into the app's seed JSON.

Usage:
    python3 tools/import_xlsx.py "path/to/workbook.xlsx" data/seed.json \
        --height <cm> --weight <kg>

The output follows docs/data-model.md (schemaVersion 1). It contains personal
data, so data/ is git-ignored; the app loads the file through its import
screen instead of shipping it in the static bundle.
"""
import argparse
import datetime as dt
import json
import os
import re
import sys
import unicodedata

import openpyxl

SCHEMA_VERSION = 1

# Canonical names for lifts that appear under several spellings.
ALIASES = {
    "back sq": "Back squat",
    "back squad": "Back squat",
    "back squat": "Back squat",
    "sq": "Back squat",
    "bench press": "Bench press",
    "dead lift": "Deadlift",
    "deadlift": "Deadlift",
    "dl": "Deadlift",
    "dead lift trap bar": "Trap bar deadlift",
    "dl trap bar": "Trap bar deadlift",
    "trap bar deadlift": "Trap bar deadlift",
    "ohp": "OHP",
    "ohp -over head press": "OHP",
    "ecc pull up": "Eccentric pull up",
    "holow body": "Hollow body",
    "hollow body": "Hollow body",
    "ortbitrek": "Orbitrek",
    "orbitrek": "Orbitrek",
    "pulldown neutral grip": "Pull down neutral grip",
    "pull down neutral grip": "Pull down neutral grip",
    "incline bench press db": "Incline bench press DB",
    "kb clean": "Kettlebell clean",
    "kettle clean": "Kettlebell clean",
    "y-raise": "Y-raise",
    "bench press close grpi": "Bench press close grip",
    "pallof press band": "Pallof press",
    "pallof press": "Pallof press",
    "sl rdl  bar": "SL RDL bar",
    "sl sq +box": "SL SQ + box",
    "sl sq + box": "SL SQ + box",
}

HEADER_FIELDS = {
    "1 rm": "oneRm",
    "nazawa ćwiczenia": "name",
    "obciążenie %": "pct",
    "[%]/[rir]": "pctRir",
    "column 10": "pctRir",
    "[kg]": "kg",
    "zrealizowane obciążenie": "actualKg",
    "reps": "reps",
    "seria": "sets",
    "tempo": "tempo",
    "odpoczynek": "rest",
    "odpoczynek [s]": "rest",
    "uwagi": "notes",
}

SECTIONS = {"rozgrzewka": "warmup", "część główna": "main", "cool down": "cooldown"}


def clean(v):
    return re.sub(r"\s+", " ", str(v)).strip() if v is not None else None


def slug(name):
    s = unicodedata.normalize("NFKD", name.lower())
    s = "".join(c for c in s if not unicodedata.combining(c)).replace("ł", "l")
    return re.sub(r"[^a-z0-9]+", "-", s).strip("-")


def canonical(name):
    n = clean(name)
    return ALIASES.get(n.lower(), n[:1].upper() + n[1:])


def num(v):
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str):
        m = re.fullmatch(r"\s*(\d+(?:[.,]\d+)?)\s*(kg)?\s*", v, re.I)
        if m:
            return float(m.group(1).replace(",", "."))
    return None


def undate(v, as_number):
    """Undo Excel turning '7.5' or '5-8' into a date (M.D order)."""
    if not isinstance(v, (dt.datetime, dt.date)):
        return v, False
    if as_number:
        return float(f"{v.month}.{v.day}"), True
    lo, hi = sorted((v.month, v.day))
    return f"{lo}-{hi}", True


def parse_rest(v):
    if isinstance(v, (int, float)):
        return int(v)
    if isinstance(v, str):
        m = re.fullmatch(r'\s*(\d+)\s*(["s]|sec)?\s*', v)
        if m:
            return int(m.group(1))
    return None


def parse_item(order, row, cols, flags_out):
    get = lambda f: row[cols[f]] if f in cols and cols[f] < len(row) else None
    item = {"order": order, "exerciseId": None, "name": canonical(get("name"))}
    flags = []

    pct, rir = None, None
    for raw in (get("pct"), get("pctRir")):
        if isinstance(raw, (int, float)) and 0 < raw <= 1.5:
            pct = round(float(raw), 3)
        elif isinstance(raw, str):
            m = re.search(r"rir\s*(\d+)", raw, re.I)
            if m:
                rir = int(m.group(1))
    item["percent1RM"] = pct
    item["rir"] = rir

    for field, key in (("kg", "plannedKg"), ("actualKg", "actualKg")):
        raw, fixed = undate(get(field), as_number=True)
        if fixed:
            flags.append(f"{key} odczytane z daty Excela")
        val = num(raw)
        item[key] = round(val, 2) if val is not None else None
        if val is None and clean(raw) not in (None, "-", "-%", ""):
            item.setdefault("loadText", clean(raw))

    reps, fixed = undate(get("reps"), as_number=False)
    if fixed:
        flags.append("reps odczytane z daty Excela")
    item["reps"] = int(reps) if isinstance(reps, float) and reps.is_integer() else (
        reps if isinstance(reps, (int, float)) else None)
    if item["reps"] is None and clean(reps) not in (None, "-"):
        item["repsText"] = clean(reps)

    sets = get("sets")
    item["sets"] = int(sets) if isinstance(sets, (int, float)) else None
    tempo = clean(get("tempo"))
    item["tempo"] = tempo if tempo and re.fullmatch(r"[\dx:]+", tempo, re.I) else None
    item["restSec"] = parse_rest(get("rest"))

    notes = [clean(get("notes"))] if get("notes") is not None else []
    extra = [clean(c) for i, c in enumerate(row) if c is not None and i > max(cols.values())]
    notes += [e for e in extra if e]
    if tempo and item["tempo"] is None:
        notes.append(tempo)
    item["notes"] = "; ".join(n for n in notes if n) or None
    if flags:
        item["importFlags"] = flags
        flags_out.extend(flags)
    return item


def parse_week_sheet(ws, flags_out):
    sessions, session, section, cols = [], None, None, None
    for row in ws.iter_rows(values_only=True):
        if not row or all(c is None for c in row):
            continue
        first = clean(row[0])
        if first and first.lower().startswith("część głowna"):
            cols = {}
            for i, h in enumerate(row):
                key = HEADER_FIELDS.get((clean(h) or "").lower())
                if key and key not in cols:
                    cols[key] = i
            session = {"order": len(sessions) + 1, "blocks": []}
            sessions.append(session)
            section = None
            continue
        if session is None:
            continue
        if first and first.lower() in SECTIONS:
            section = {"type": SECTIONS[first.lower()], "items": []}
            session["blocks"].append(section)
            continue
        name_idx = cols.get("name", 2)
        if section is None or name_idx >= len(row) or row[name_idx] is None:
            continue
        order = int(row[0]) if isinstance(row[0], (int, float)) else len(section["items"]) + 1
        section["items"].append(parse_item(order, row, cols, flags_out))
    return [s for s in sessions if any(b["items"] for b in s["blocks"])]


def classify_week(sheet):
    n = sheet.strip()
    for pattern, program in (
        (r"^(\d+) Tydzień SOLO$", "solo"),
        (r"^(\d+) Tydzień 2\.0$", "blok-2"),
        (r"^(\d+) (?:Tydzień|Competition Week)", "blok-1"),
        (r"^Tydzień (\d+)$", "blok-1"),
    ):
        m = re.match(pattern, n)
        if m:
            label = None
            if "Taper" in n:
                label = "Taper"
            elif "Competition" in n:
                label = "Zawody"
            return program, int(m.group(1)), label
    return None


PROGRAMS = {
    "blok-1": "Blok 1 (z trenerem)",
    "blok-2": "Blok 2.0",
    "solo": "SOLO",
}


def parse_max(ws):
    rows = [r for r in ws.iter_rows(values_only=True) if r and any(c is not None for c in r)]
    header = rows[0]
    later = [(i, c) for i, c in enumerate(header) if isinstance(c, dt.datetime)]
    out = []
    for r in rows[1:]:
        name = canonical(r[0])
        if r[1] is not None:
            out.append({
                "exercise": name, "date": None, "weightKg": num(r[1]),
                "reps": int(r[2]) if isinstance(r[2], (int, float)) else None,
                "estimated1RM": num(r[3]) if len(r) > 3 else None,
                "source": "MAX (data nieznana)",
            })
        for i, d in later:
            if i < len(r) and num(r[i]) is not None:
                out.append({
                    "exercise": name, "date": d.date().isoformat(), "weightKg": num(r[i]),
                    "reps": 1, "estimated1RM": num(r[i]), "source": "MAX",
                })
    return out


def parse_daily(ws):
    logs, dates = [], None
    for r in ws.iter_rows(values_only=True):
        if not r or r[0] is None:
            continue
        label = clean(r[0]).lower()
        vals = list(r[1:8])
        if label.startswith("data"):
            dates = [d.date().isoformat() if isinstance(d, dt.datetime) else None for d in vals]
            continue
        if not dates:
            continue
        for i, d in enumerate(dates):
            if d is None or i >= len(vals) or vals[i] in (None, 0):
                continue
            entry = next((l for l in logs if l["date"] == d), None)
            if entry is None:
                entry = {"date": d}
                logs.append(entry)
            v = vals[i]
            if label.startswith("ile kroków"):
                entry["steps"] = int(v)
            elif label.startswith("ile godzin snu"):
                secs = v.total_seconds() if isinstance(v, dt.timedelta) else (
                    v.hour * 3600 + v.minute * 60 if isinstance(v, dt.time) else None)
                if secs:
                    entry["sleepMin"] = int(secs // 60)
            elif label.startswith("czy był trening"):
                entry["training"] = clean(v)
            elif label.startswith("ile kalorii"):
                entry["kcal"] = int(v)
    return sorted(logs, key=lambda l: l["date"])


def parse_pt_cards(wb):
    cards = []
    for name in wb.sheetnames:
        if "Karta treningowa" not in name:
            continue
        m = re.match(r"^(\d+)\.0", name)
        card = {"card": int(m.group(1)) if m else 1, "sessionDates": []}
        for r in wb[name].iter_rows(values_only=True):
            if not r or r[0] is None:
                continue
            if isinstance(r[0], (int, float)) and len(r) > 1 and isinstance(r[1], dt.datetime):
                card["sessionDates"].append(r[1].date().isoformat())
            elif clean(r[0]).startswith("Ilość") or "Ilość wykupionych treningów:" in [clean(c) for c in r]:
                pass
        cards.append(card)
    return sorted(cards, key=lambda c: c["card"])


def parse_warmups(ws):
    templates, cur = [], None
    for r in ws.iter_rows(values_only=True):
        if not r or all(c is None for c in r):
            continue
        if isinstance(r[0], str):
            cur = {"name": clean(r[0]), "items": []}
            templates.append(cur)
        elif cur is not None:
            cells = [c for c in r[1:] if c is not None]
            if not cells:
                continue
            rest = [clean(c) for c in cells[1:]]
            cur["items"].append({
                "order": int(r[0]), "name": clean(cells[0]),
                "dose": rest[0] if rest else None,
                "sets": int(float(rest[1])) if len(rest) > 1 else None,
            })
    return templates


def parse_goals(ws):
    rows = [[clean(c) if not isinstance(c, (int, float)) else c for c in r]
            for r in ws.iter_rows(values_only=True) if r and any(c is not None for c in r)]
    goals, details, legacy = [], {}, {}
    for r in rows:
        r = [c for c in r if c is not None]
        head = str(r[0])
        if head.lower().startswith("zapotrzebowanie"):
            legacy["maintenanceKcal"] = r[1]
        elif head.lower().startswith("aby redukować"):
            legacy["reductionKcal"] = r[1]
        elif head in ("Białko", "Tłuszcze", "Węglowodany"):
            legacy[{"Białko": "proteinG", "Tłuszcze": "fatG", "Węglowodany": "carbsG"}[head]] = round(r[3])
        elif head != "CELE":
            goals.append(head)
            items = [c for c in r[1:] if isinstance(c, str) and c not in ("wiek", "wzrost", "waga", "lat", "cm", "kg")]
            if items:
                details[head] = items
    return {"focus": goals, "details": details}, legacy


def parse_diet(ws):
    rows = list(ws.iter_rows(values_only=True))
    header = [clean(c) for c in rows[0]]
    cols = {}
    for i, h in enumerate(header):
        if h and h not in cols:
            cols[h] = []
    for r in rows[1:]:
        for i, c in enumerate(r or []):
            if i < len(header) and header[i] in cols and c is not None:
                cols[header[i]].append(clean(c))
    # The sheet keeps sub-headings (ending with ':') inside columns.
    groups = {}
    for h, items in cols.items():
        key = h.rstrip(":")
        groups.setdefault(key, [])
        for it in items:
            if it.endswith(":") or ("(" in it and it.endswith("):")):
                key = it.rstrip(":")
                groups.setdefault(key, [])
            else:
                if it not in groups[key]:
                    groups[key].append(it)
    return [{"group": k, "items": v} for k, v in groups.items() if v]


def parse_tips(ws):
    tips, cur = [], None
    for r in ws.iter_rows(values_only=True):
        cells = [clean(c) for c in (r or []) if c is not None]
        if not cells:
            continue
        if len(cells) == 1 and cells[0].endswith(":") or cells[0].startswith("Wskazówki"):
            cur = {"section": cells[0].rstrip(":."), "items": []}
            tips.append(cur)
        elif cur is not None:
            cur["items"].append(cells[-1])
    return tips


def prune(v):
    """Drop empty fields so the seed stays small and readable."""
    if isinstance(v, dict):
        return {k: prune(x) for k, x in v.items() if x is not None or k in KEEP_NULL}
    if isinstance(v, list):
        return [prune(x) for x in v]
    return v


KEEP_NULL = {"heightCm", "weightKg", "birthYear", "sex", "activityLevel", "date", "label", "exerciseId"}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("out")
    ap.add_argument("--height", type=float)
    ap.add_argument("--weight", type=float)
    a = ap.parse_args()

    wb = openpyxl.load_workbook(a.xlsx, read_only=True, data_only=True)
    flags = []
    programs = {k: {"id": k, "name": v, "weeks": []} for k, v in PROGRAMS.items()}
    for sheet in wb.sheetnames:
        c = classify_week(sheet)
        if not c:
            continue
        program, number, label = c
        sessions = parse_week_sheet(wb[sheet], flags)
        programs[program]["weeks"].append({
            "id": f"{program}-w{number}", "number": number, "label": label,
            "sourceSheet": sheet.strip(), "sessions": sessions,
        })
    for p in programs.values():
        p["weeks"].sort(key=lambda w: w["number"])

    exercises = {}
    for p in programs.values():
        for w in p["weeks"]:
            for s in w["sessions"]:
                for b in s["blocks"]:
                    for it in b["items"]:
                        ex = exercises.setdefault(slug(it["name"]), {"id": slug(it["name"]), "name": it["name"], "uses": 0})
                        ex["uses"] += 1
                        it["exerciseId"] = ex["id"]
                        del it["name"]

    maxes = parse_max(wb["MAX"])
    for m in maxes:
        ex = exercises.setdefault(slug(m["exercise"]), {"id": slug(m["exercise"]), "name": m["exercise"], "uses": 0})
        m["exerciseId"] = ex["id"]
        del m["exercise"]

    goals, legacy = parse_goals(wb["Cele"])
    seed = {
        "schemaVersion": SCHEMA_VERSION,
        "generatedAt": dt.datetime.now(dt.timezone.utc).isoformat(timespec="seconds"),
        "source": os.path.basename(a.xlsx),
        "profile": {
            "heightCm": a.height, "weightKg": a.weight,
            "birthYear": None, "sex": None, "activityLevel": None,
        },
        "goals": goals,
        "nutrition": {
            "legacyTargets": dict(legacy, note="Z arkusza, liczone dla wagi z arkusza; do przeliczenia"),
            "products": parse_diet(wb["Dieta produkty"]),
        },
        "tips": parse_tips(wb["Wskazówki"]),
        "warmupTemplates": parse_warmups(wb["Rozgrzewka"]),
        "exercises": sorted(exercises.values(), key=lambda e: e["name"].lower()),
        "oneRepMaxes": maxes,
        "programs": [programs[k] for k in ("blok-1", "blok-2", "solo")],
        "dailyLogs": parse_daily(wb["Pomiary "]),
        "bodyMeasurements": [],
        "ptSessions": parse_pt_cards(wb),
        "workoutLogs": [],
    }
    with open(a.out, "w", encoding="utf-8") as f:
        json.dump(prune(seed), f, ensure_ascii=False, indent=1)

    n_sessions = sum(len(w["sessions"]) for p in seed["programs"] for w in p["weeks"])
    n_items = sum(len(b["items"]) for p in seed["programs"] for w in p["weeks"]
                  for s in w["sessions"] for b in s["blocks"])
    print(f"programs={len(seed['programs'])} weeks={sum(len(p['weeks']) for p in seed['programs'])} "
          f"sessions={n_sessions} items={n_items} exercises={len(exercises)} "
          f"1RM={len(maxes)} dailyLogs={len(seed['dailyLogs'])} flags={len(flags)}", file=sys.stderr)


if __name__ == "__main__":
    main()
